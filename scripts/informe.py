"""Genera el informe del proyecto en Word (docs/informe_proyecto.docx) y, con Word, en PDF.

Parte I: la base de datos. Parte II: estadística descriptiva. Parte III: cómo se construyó todo
(código, librerías, mapa, R, web, publicación). Anexos: diccionario de variables y estructura del repo.

Uso: python scripts/informe.py          (sin -I: importa TIPOS de datos_dashboard.py)
La conversión a PDF y la actualización del índice usan Microsoft Word por COM (PowerShell); si Word no
está, queda solo el .docx.
"""
import json
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from datos_dashboard import TIPOS

RAIZ = Path(__file__).resolve().parent.parent
DOCX = RAIZ / "docs" / "informe_proyecto.docx"
PDF = RAIZ / "docs" / "informe_proyecto.pdf"
WEB = "https://crisrumiguano-art.github.io/geoquimica-earthchem-ml/"
REPO = "https://github.com/crisrumiguano-art/geoquimica-earthchem-ml"
AZUL = RGBColor(0x1C, 0x5C, 0xAB)
GRIS = RGBColor(0x52, 0x51, 0x4E)


def miles(n):
    return f"{n:,}".replace(",", ".")


# ------------------------------------------------------------------ utilidades de formato

def campo(parrafo, instruccion):
    """Inserta un campo de Word (PAGE, TOC…); Word lo calcula al abrir o al actualizar."""
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, instruccion), ("separate", None), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = texto
        run._r.append(el)


def sombrear(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tc.append(shd)


class Informe:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21)
        for lado in ("left_margin", "right_margin"):
            setattr(s, lado, Cm(2.3))
        s.top_margin = s.bottom_margin = Cm(2.2)
        normal = self.d.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for nivel, tam in ((1, 17), (2, 13.5), (3, 11.5)):
            h = self.d.styles[f"Heading {nivel}"]
            h.font.name, h.font.size, h.font.color.rgb = "Calibri", Pt(tam), AZUL
        pie = s.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = pie.add_run("Geoquímica de rocas · Estadística y Probabilidades · FIGEMPA-UCE · pág. ")
        r.font.size, r.font.color.rgb = Pt(8), GRIS
        campo(pie, "PAGE")

    def h(self, texto, nivel=1):
        return self.d.add_heading(texto, nivel)

    def p(self, *partes, estilo=None):
        """partes: str o tuplas (texto, 'b'|'i'|'c') para negrita, cursiva o código en línea."""
        par = self.d.add_paragraph(style=estilo)
        for parte in partes:
            texto, fmt = (parte, "") if isinstance(parte, str) else parte
            run = par.add_run(texto)
            run.bold, run.italic = "b" in fmt, "i" in fmt
            if "c" in fmt:
                run.font.name, run.font.size = "Consolas", Pt(9.5)
        return par

    def lista(self, items, numerada=False):
        for it in items:
            self.p(*(it if isinstance(it, tuple) else (it,)), estilo="List Number" if numerada else "List Bullet")

    def codigo(self, texto):
        t = self.d.add_table(rows=1, cols=1)
        c = t.rows[0].cells[0]
        sombrear(c, "F2F1EC")
        c.paragraphs[0].text = ""
        for i, linea in enumerate(texto.strip("\n").split("\n")):
            par = c.paragraphs[0] if i == 0 else c.add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            run = par.add_run(linea)
            run.font.name, run.font.size = "Consolas", Pt(8.5)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def tabla(self, filas, anchos=None):
        t = self.d.add_table(rows=len(filas), cols=len(filas[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, fila in enumerate(filas):
            for j, valor in enumerate(fila):
                c = t.cell(i, j)
                c.text = ""
                run = c.paragraphs[0].add_run(str(valor))
                run.font.size = Pt(9)
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                if i == 0:
                    run.bold = True
                    sombrear(c, "E8F0FB")
                if anchos:
                    c.width = Cm(anchos[j])
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def imagen(self, ruta, ancho_cm, pie):
        self.d.add_picture(str(ruta), width=Cm(ancho_cm))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        par = self.p((pie, "i"))
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.runs[0].font.size, par.runs[0].font.color.rgb = Pt(9), GRIS

    def salto(self):
        self.d.add_page_break()


# ------------------------------------------------------------------ contenido

def portada(I):
    for _ in range(5):
        I.d.add_paragraph()
    t = I.p(("Geoquímica de rocas a escala mundial", "b"))
    t.runs[0].font.size, t.runs[0].font.color.rgb = Pt(28), AZUL
    s = I.p("Estadística y machine learning sobre 159.760 análisis de roca total de EarthChem PetDB")
    s.runs[0].font.size, s.runs[0].font.color.rgb = Pt(14), GRIS
    I.p(("Informe del proyecto: la base de datos, la estadística descriptiva y la documentación técnica de "
         "cómo se construyó todo, incluida la página web.", "i"))
    for _ in range(3):
        I.d.add_paragraph()
    I.p(("Universidad Central del Ecuador", "b"))
    I.p("Facultad de Ingeniería en Geología, Minas, Petróleos y Ambiental (FIGEMPA) · Carrera de Ingeniería en Minas")
    I.p("Asignatura: Estadística y Probabilidades · Docente: Dr. Christian Mejía E., Ph.D.")
    I.d.add_paragraph()
    I.p(("Integrantes", "b"))
    for n in ("Cristoffer Joel Rumiguano Chariguaman", "Ana Gabriela Fernández Chuquimarca",
              "Evelyn Maritza Bustamante Armijos", "Jessica Jaqueline Haro Simba"):
        I.p(n)
    I.d.add_paragraph()
    I.p("Sitio web: ", (WEB, "c"))
    I.p("Repositorio: ", (REPO, "c"))
    I.p("Quito, octubre de 2026")
    I.salto()
    c = I.p(("Contenido", "b"))  # no es título: si no, el índice se listaría a sí mismo
    c.runs[0].font.size, c.runs[0].font.color.rgb = Pt(17), AZUL
    campo(I.d.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')


def guia_equipo(I, df):
    # Salto en el título, no como párrafo aparte: si el índice llena la página, un salto suelto deja una hoja en blanco.
    I.h("Guía rápida para el equipo", 1).paragraph_format.page_break_before = True
    I.p("Esta guía resume, en pocas páginas, qué está hecho, dónde está cada cosa y qué hay que saber para la exposición. "
        "El resto del informe explica todo con detalle.")
    I.h("Dónde está cada cosa", 2)
    I.tabla([
        ["Qué", "Dónde"],
        ["Página web del proyecto", WEB],
        ["Estadística descriptiva (tablas y gráficas)", WEB + "descriptiva/"],
        ["Tabla de variables (dominio, rango, etc.)", WEB + "descriptiva/#variables  ·  data/tabla_variables.xlsx"],
        ["Este informe (PDF y Word)", WEB + "informe_proyecto.pdf  ·  " + WEB + "informe_proyecto.docx"],
        ["Base limpia (159.760 × 75)", "data/earthchem_limpio.csv.gz en el repositorio"],
        ["Mapa en QGIS", "data/gis/mapa_ubicacion.qgz (abrir junto con mapa_ubicacion.gpkg)"],
        ["Código en R", "carpeta R/ del repositorio"],
        ["Repositorio (todo el código)", REPO],
        ["Datos originales de EarthChem (~600 MB)", "Google Drive de Dalember: TRABAJOS ESTADISTICA › 10-26 GDRIVE "
                                                   "GEOQUIMICA ROCAS › datos_crudos_earthchem (pedir acceso)"],
    ], [6, 10.3])
    I.h("Qué tiene la página web", 2)
    I.lista([
        (("Inicio: ", "b"), "introducción, planteamiento del problema, mapa de ubicación (QGIS, Python e interactivo) y objetivos."),
        (("Descriptiva (Unidad 1): ", "b"), "población, individuo y muestra; tabla de variables; tablas de frecuencia y gráficas "
         "de las variables cualitativas, discretas y continuas; tratamiento de outliers; gráficos interactivos."),
        (("Datos: ", "b"), "de dónde salieron los datos, cómo se prepararon y todas las descargas."),
        (("Equipo: ", "b"), "integrantes y docente."),
        (("Inferencial, Regresión y Machine learning: ", "b"), "aparecen como «pronto»; se completan en las próximas unidades."),
    ])
    I.h("Estado por lección", 2)
    I.tabla([
        ["Lección", "Qué pide", "Estado"],
        ["L1–L2", "introducción, problema, objetivos; población, individuo y muestra", "Hecho (web: Inicio y Descriptiva 01)"],
        ["L3", "muestreo y mapa de ubicación en QGIS", "Mapa hecho (QGIS y Python). Falta el informe de muestreo estratificado"],
        ["L4", "tabla de variables: tipo, subtipo, dominio, rango, instrumento, unidad", "Hecho (web y Excel, formato del curso)"],
        ["L5", "tablas de frecuencia (TDF) y gráficas", "Hecho: 4 cualitativas, 2 discretas, 3 continuas (SiO2, MgO, Cr)"],
        ["L6", "indicadores de posición, dispersión y forma", "Hecho para esas variables; se puede ampliar a más"],
        ["L7", "exposición de la etapa descriptiva y conclusiones", "Pendiente: redactar conclusiones y preparar la exposición"],
        ["Unidades 2–4", "inferencia, regresión, red neuronal", "Pendiente"],
    ], [2.2, 7, 7.1])
    I.h("Cinco ideas clave para la exposición", 2)
    I.lista([
        (("La muestra no es aleatoria. ", "b"), "PetDB reúne lo que los investigadores decidieron estudiar durante casi un siglo, "
         "con más rocas del fondo del mar que de los continentes. Para la inferencia se usarán submuestras aleatorias por estrato."),
        (("Una celda vacía significa «no se midió», no cero. ", "b"), "Por eso cada análisis usa solo las filas que tienen "
         "medidas las variables de ese análisis, y siempre se informa cuántas son (n)."),
        (("La agrupación por defecto no servía. ", "b"), "Con Sturges sobre el rango completo, el 99,9 % del cromo caía en un solo "
         "intervalo. Se usaron intervalos sobre el rango típico y dos clases abiertas para los extremos."),
        (("Los outliers no se borraron. ", "b"), "Casi todos son otras rocas (riolitas, peridotitas, cromititas), no errores. "
         "Borrarlos habría eliminado poblaciones enteras."),
        (("Los elementos traza se miran en escala logarítmica. ", "b"), "Así se ve su forma real y aparecen dos poblaciones en el "
         "cromo: basaltos y rocas ultramáficas."),
    ], numerada=True)
    I.h("Cómo trabajar con los archivos en tu computadora", 2)
    I.p(("Ver la base en Excel. ", "b"), "Descarga data/earthchem_limpio.csv.gz desde la página Datos, descomprímela con 7-Zip y "
        "en Excel usa Datos › Desde texto/CSV, con origen UTF-8 y, en Transformar datos, configuración regional Inglés "
        "(Estados Unidos) para que el punto sea el separador decimal.")
    I.p(("Correr los scripts de R. ", "b"), "Instala R (4.4 o más nuevo) y RStudio. En GitHub, en el repositorio, usa Code › "
        "Download ZIP y descomprime. Abre RStudio, ve a Session › Set Working Directory y elige la carpeta del repositorio. "
        "La primera vez ejecuta install.packages(c(\"e1071\", \"openxlsx\")). Después, por ejemplo, source(\"R/03_continuas.R\"): "
        "las gráficas quedan en docs/img/r/ y las tablas en R/salidas/.")
    I.p(("Abrir el mapa en QGIS. ", "b"), "Descarga mapa_ubicacion.qgz y mapa_ubicacion.gpkg (página Datos) en la misma carpeta "
        "y abre el .qgz. La composición de impresión se llama «Mapa de ubicación A4».")
    I.h("Glosario", 2)
    I.tabla([
        ["Término", "Qué significa"],
        ["TDF", "tabla de distribución de frecuencias: cuántos datos caen en cada clase o intervalo"],
        ["Li, Ls, MC", "límite inferior y superior del intervalo, y su marca de clase (el punto medio)"],
        ["ni, hi", "frecuencia absoluta (cuántos datos) y relativa (qué porcentaje)"],
        ["Ni, Hi ↑ / ↓", "frecuencias acumuladas ascendentes y descendentes; con ellas se dibuja la ojiva"],
        ["Sturges", "regla para el número de intervalos: k = 1 + 3,3 log n"],
        ["IQR (rango intercuartílico)", "Q3 − Q1; un dato es atípico si cae fuera de Q1 − 1,5·IQR o de Q3 + 1,5·IQR"],
        ["Clase abierta", "intervalo sin uno de sus límites («< 36» o «≥ 66») que agrupa los valores extremos"],
        ["Error de agrupación", "diferencia entre la media (o desviación) calculada con la TDF y con los datos sin agrupar"],
        ["Escala log10", "trabajar con el logaritmo del valor; separa bien datos que van de 0,01 a 100.000"],
        ["Estrato", "subgrupo homogéneo de la muestra (por ejemplo, solo rocas volcánicas máficas)"],
        ["Dominio / rango", "valores que la variable puede tomar en teoría / valores que tomó en la muestra (ver sección 4.1)"],
        ["wt% / ppm", "porcentaje en peso / partes por millón (1 wt% = 10.000 ppm)"],
    ], [4.5, 11.8])
    I.salto()


def parte_base(I, df):
    n, ncol = df.shape
    mayores = ["SiO2 (wt%)", "Al2O3 (wt%)", "MgO (wt%)", "CaO (wt%)", "Na2O (wt%)", "K2O (wt%)"]
    completos = int(df[mayores].notna().all(axis=1).sum())
    articulos = df["Citation"].nunique()
    oceano = df["Geographic Features"].fillna("").str.contains("OCEAN").mean()
    tec = df["Tectonic Setting"].value_counts()
    sub = df["Sample SubType"].value_counts()
    reciente = (df["Publication Year"] >= 2000).mean()

    I.h("Parte I. La base de datos", 1)
    I.h("1. Qué es esta base", 2)
    I.p(f"La base reúne {miles(n)} análisis químicos de rocas tomados de muestras recolectadas en todo el mundo. Cada "
        "fila corresponde al análisis de una muestra de ", ("roca total", "i"), ", es decir, de la roca molida y medida en "
        "conjunto, no de un mineral en particular. Para cada muestra se informa su composición en elementos mayores "
        "(sílice, aluminio, magnesio, calcio, sodio, potasio y otros, expresados como óxidos en porcentaje en peso), en "
        "elementos traza (en partes por millón) y en tierras raras, junto con el tipo de roca, el lugar donde se tomó, "
        "el ambiente geológico en que se formó y el artículo científico que publicó el dato.")
    I.p(f"Predominan las rocas ígneas, y en particular los basaltos ({miles(sub.get('BASALT', 0))} análisis), seguidos "
        f"por peridotitas, tefras, harzburgitas, diabasas, lherzolitas, gabros y andesitas. Cerca del {oceano:.0%} de las "
        "muestras proviene del fondo oceánico: PetDB nació para reunir los datos de las dorsales oceánicas y de las "
        "perforaciones de los programas internacionales de exploración del océano (DSDP, ODP e IODP).")

    I.h("2. De dónde salieron los datos", 2)
    I.p("Los datos provienen de ", ("PetDB", "b"), " (Petrological Database), una de las bases que integran ",
        ("EarthChem", "b"), ", la infraestructura de datos geoquímicos financiada por la National Science Foundation de "
        "Estados Unidos. Hoy PetDB se consulta a través de la Astromaterials Data Synthesis (search.astromat.org) y del "
        "portal de EarthChem (www.earthchem.org/petdb).")
    I.p(f"PetDB no genera datos propios: su equipo los transcribe de artículos científicos revisados por pares y los "
        f"ordena en un formato común. La exportación usada reúne datos de {miles(articulos)} artículos publicados entre "
        f"{int(df['Publication Year'].min())} y {int(df['Publication Year'].max())}; algo más de la mitad ({reciente:.0%}) "
        "corresponde a trabajos publicados desde el año 2000. Cada fila conserva la cita del artículo de origen y un enlace "
        "a la ficha de la muestra en PetDB, de modo que cualquier valor se puede rastrear hasta su publicación.")
    I.p("La descarga se hizo el ", ("13 de octubre de 2026", "b"), " con un único filtro: material analizado igual a roca "
        "total (incluido el registrado como roca total o vidrio sin especificar). Se pidieron todas las familias de "
        "variables disponibles: elementos mayores, trazas, tierras raras, edades, miembros extremos, moda mineral, gases "
        "nobles y volátiles. La consulta exacta quedó guardada en ", ("data/raw/earthchem_export_oct13_26_142339.txt", "c"),
        ", así que la descarga se puede repetir. Licencia ", ("Creative Commons Atribución 4.0 (CC BY 4.0)", "b"),
        ": los datos pueden usarse y redistribuirse citando la fuente.")

    I.h("3. Cómo se preparó", 2)
    I.p("La exportación original venía en dos archivos que suman unos 600 MB: uno con los resultados químicos, de 2.318 "
        "columnas, y otro con los datos descriptivos de cada muestra. Las columnas son tantas porque cada variable química "
        "viene acompañada de tres columnas de texto (método analítico, comentario y laboratorio). La preparación consistió en:")
    I.lista([
        "unir los dos archivos, verificando que las muestras coincidieran fila por fila;",
        "quitar las columnas de método, comentario y laboratorio;",
        "conservar las variables químicas con datos en al menos el 5 % de las muestras (quedan 52; se descartan, entre "
        "otras, las relaciones isotópicas, presentes en menos del 2 % de las filas);",
        "convertir las coordenadas a grados decimales, con el sur y el oeste en negativo;",
        "dejar vacía, y no en cero, toda celda química sin medición, y también las concentraciones negativas, que son "
        "lecturas bajo el límite de detección (salvo la pérdida por calcinación, que puede ser negativa de verdad);",
        "corregir cinco muestras de la dorsal de Gakkel que venían con latitud y longitud invertidas;",
        "construir cinco variables nuevas: año de publicación, número de analitos, clase de sílice, grado de alteración "
        "y ambiente tectónico agrupado.",
    ], numerada=True)

    I.h("4. Variables", 2)
    I.p(f"La base final tiene {miles(n)} filas y {ncol} columnas, con todos los tipos de variable que se usan en estadística:")
    I.tabla([
        ["Tipo", "Variables", "Comentario"],
        ["Nominales", "Sample Type, Sample SubType, Tectonic Setting, Tectonic Group, Geographic Features, Rock Texture, "
                      "Expedition, Citation",
         f"Categorías sin orden. El ambiente tectónico tiene {df['Tectonic Setting'].nunique()} valores; Tectonic Group "
         "los resume en tres grandes ambientes más los mixtos."],
        ["Nominal binaria", "Analyzed Material", "Dos valores: roca total y roca total o vidrio sin especificar."],
        ["Ordinales", "SiO2 Class, Alteration Grade, Geologic Age, Sample Alteration",
         "SiO2 Class: ultrabásica (< 45 %), básica (45–52 %), intermedia (52–63 %), ácida (≥ 63 %). Alteration Grade: "
         "0 (fresca) a 5 (totalmente alterada)."],
        ["Numéricas discretas", "Publication Year, N Analytes", "Solo valores enteros."],
        ["Numéricas continuas", "52 químicas, Latitude, Longitude, elevaciones, edades (Ma)",
         "Óxidos mayores en % en peso; trazas y tierras raras en ppm."],
        ["Identificadores", "Sample Name, Sample URL", "No se analizan; sirven para rastrear cada muestra."],
    ], [3.2, 6.3, 7])
    I.p("La tabla de variables completa sigue el formato de la plantilla del curso: nombre, descripción, tipo, subtipo, "
        "dominio (como definición y como conjunto, por ejemplo D={x/x es concentración ^ x ∈ R} y D={x/x ∈ R+ ∪ {0}}), "
        "rango observado (por ejemplo R={y/ 0.0011 ≤ y ≤ 98.9} para el SiO2), instrumento analítico, unidad y escala de "
        "medición. Está en la página web (Estadística descriptiva) y en ", ("data/tabla_variables.xlsx", "c"),
        ". El diccionario resumido está en el Anexo A.")

    I.h("4.1 Dominio y rango: qué significan", 3)
    I.p("La idea es la misma que con las funciones en matemáticas. El ", ("dominio", "b"), " son todos los valores que la "
        "variable podría tomar en teoría, por cómo es la variable. El ", ("rango", "b"), " son los valores que tomó de verdad en "
        "nuestros datos. La plantilla del curso escribe el dominio dos veces: en palabras (la definición) y como conjunto "
        "numérico. Por ejemplo, una concentración puede ser cualquier número real desde cero hacia arriba (dominio R+ ∪ {0}), "
        "pero en esta base el SiO2 va de 0,0011 % a 98,9 % (rango).")
    vars_ = {v["Variable"]: v for v in json.loads((RAIZ / "docs" / "data" / "variables.json").read_text(encoding="utf-8"))}
    I.tabla([["Variable", "Dominio (definición)", "Dominio (conjunto)", "Rango"]] + [
        [c, vars_[c]["Dominio (definición)"], vars_[c]["Dominio (conjunto)"], vars_[c]["Rango"]]
        for c in ("SiO2 (wt%)", "Publication Year", "Alteration Grade", "Sample Type") if c in vars_
    ], [3, 5.6, 3.4, 4.3])
    I.p("En las variables cualitativas el rango son las categorías que aparecen en la base; cuando son pocas se enumeran y "
        "cuando son muchas se indica cuántas hay.")

    I.h("5. Requisitos del trabajo", 2)
    I.tabla([["Requisito", "Exigido", "Base"], ["Filas", "al menos 10.000", miles(n)], ["Columnas", "al menos 30", str(ncol)],
             ["Tipos de variables", "nominales, ordinales, discretas y continuas", "todos presentes"]], [5, 6, 5.3])

    I.h("6. Limitaciones que conviene conocer", 2)
    I.p(("No todas las muestras tienen todas las mediciones. ", "b"), f"Cada artículo midió lo que necesitaba, así que la "
        f"base tiene muchos huecos: solo {miles(completos)} filas ({completos / n:.0%}) tienen completos los seis óxidos "
        "principales, y las tierras raras aparecen en una de cada cinco filas. Una celda vacía significa que el valor no "
        "se midió; conviene filtrar las filas completas para cada análisis en lugar de rellenar huecos sin criterio.")
    I.p(("Falta el TiO2. ", "b"), "El óxido de titanio prácticamente no vino en la exportación (cuatro valores). Si el "
        "análisis lo requiere, habrá que volver a descargarlo desde PetDB.")
    I.p(("Unidades y formas distintas. ", "b"), "Níquel, cobalto, cobre y zinc aparecen en % en peso y en ppm (1 % = "
        "10.000 ppm), y el hierro se informa de cuatro maneras (FeO, Fe2O3, FeOT y Fe2O3T).")
    I.p(("Sesgos de origen. ", "b"), f"Hay muchos más basaltos de fondo oceánico que rocas continentales y los métodos de "
        f"laboratorio cambiaron en casi un siglo. El ambiente más frecuente, {tec.index[0].lower()}, reúne "
        f"{miles(int(tec.iloc[0]))} muestras, mientras otros tienen apenas decenas. La muestra no es probabilística: la "
        "inferencia deberá hacerse con submuestras aleatorias estratificadas.")

    I.h("7. Acceso y cita", 2)
    I.p("La base limpia está en ", ("data/earthchem_limpio.csv.gz", "c"), " (CSV comprimido, 26 MB). Para abrirla en Excel: "
        "descomprimir (por ejemplo con 7-Zip) e importar desde Datos › Desde texto/CSV con codificación UTF-8 y configuración "
        "regional Inglés (Estados Unidos), para que el punto se lea como separador decimal. En R o Python se lee directo.")
    I.p("Cita sugerida: ", ("Datos descargados de EarthChem PetDB / Astromaterials Data Synthesis (https://search.astromat.org/) "
        "el 13 de octubre de 2026; material analizado = roca total. Licencia CC BY 4.0.", "i"))
    I.salto()


def parte_descriptiva(I, df):
    I.h("Parte II. Estadística descriptiva", 1)
    I.p("La estadística se calcula en R, como exige la asignatura, siguiendo la estructura de las plantillas del curso: "
        "tablas de distribución de frecuencias (TDF) con límites, marca de clase, frecuencias absolutas y relativas y sus "
        "acumuladas ascendentes y descendentes; histograma con polígono de frecuencias, ojiva y diagrama de caja; e "
        "indicadores de posición, dispersión y forma. Los resultados completos están en la página web.")
    I.h("8. Por qué no sirvió la agrupación por defecto", 2)
    I.p("Las plantillas calculan k por Sturges (k = 1 + 3,3 log n) y la amplitud como a = R / k. En esta base el recorrido "
        "está dominado por unos pocos valores extremos: el cromo va de 0,002 a 295.000 ppm, cada intervalo mide 18.437 ppm "
        "y el primero se queda con el 99,9 % de los datos. Las reglas √n y 2√n dan más de 190 intervalos.")
    I.imagen(RAIZ / "docs" / "img" / "r" / "Cr_agrupacion.png", 16, "Cromo: agrupación por defecto (izquierda) frente a la corregida en log10 (derecha).")
    I.h("9. Outliers: otras rocas, no errores", 2)
    I.p("La regla intercuartílica marca entre el 8 y el 17 % de los datos, y casi todos son rocas reales: tefras y rocas "
        "ácidas en la sílice, peridotitas en el magnesio, cromititas en el cromo. La base mezcla poblaciones, y la regla "
        "supone una sola. Por eso no se elimina ningún dato. La TDF corregida usa intervalos de Sturges sobre el rango "
        "típico (los límites intercuartílicos), con amplitud redonda, y dos clases abiertas para los extremos. El error de "
        "agrupación (media y desviación agrupadas frente a datos totales) queda por debajo del 5 % en todos los casos. Los "
        "elementos traza se analizan en log10 y todo el análisis se repite en un estrato homogéneo (rocas volcánicas máficas).")
    I.imagen(RAIZ / "docs" / "img" / "r" / "SiO2_toda_la_muestra.png", 16, "SiO2: histograma con polígono, ojiva y diagrama de caja de la TDF corregida.")
    I.p("El razonamiento completo, con todas las cifras, está en ", ("R/analisis_agrupacion_outliers.md", "c"), ".")
    I.h("10. ¿Alcanzan los datos para las regresiones?", 2)
    I.p("Sí. Aunque hay muchos huecos, en una regresión solo cuentan las filas que tienen medidas las dos variables, y "
        "quedan miles. La tabla muestra algunos pares útiles para la Unidad 3, con el número de filas completas y la "
        "correlación (r), en escala original y en escala logarítmica:")
    d = df.copy()
    d["Na2O+K2O"] = d["Na2O (wt%)"] + d["K2O (wt%)"]
    filas = [["Variable X", "Variable Y", "Filas completas", "r", "r (log10)"]]
    for x, y in [("MgO (wt%)", "Ni (ppm)"), ("La (ppm)", "Ce (ppm)"), ("Zr (ppm)", "Hf (ppm)"), ("Nb (ppm)", "Ta (ppm)"),
                 ("Rb (ppm)", "K2O (wt%)"), ("Th (ppm)", "U (ppm)"), ("MgO (wt%)", "Cr (ppm)"), ("SiO2 (wt%)", "Na2O+K2O")]:
        p = d[[x, y]].dropna()
        p = p[(p[x] > 0) & (p[y] > 0)]
        r = p[x].corr(p[y])
        rl = np.log10(p[x]).corr(np.log10(p[y]))
        filas.append([x, y, miles(len(p)), f"{r:.2f}".replace(".", ","), f"{rl:.2f}".replace(".", ",")])
    I.tabla(filas, [3.3, 3.3, 3, 2.2, 2.2])
    mayores = ["SiO2 (wt%)", "Al2O3 (wt%)", "MgO (wt%)", "CaO (wt%)", "Na2O (wt%)", "K2O (wt%)"]
    I.p(f"Para una regresión múltiple con los seis óxidos mayores quedan {miles(int(df[mayores].notna().all(axis=1).sum()))} "
        "filas completas. Tres precauciones:")
    I.lista([
        (("Usar escala logarítmica en los elementos traza. ", "b"), "Zr y Hf parecen no relacionarse en escala normal (r bajo), "
         "pero en log10 la relación es casi perfecta: unos pocos valores extremos arruinan la regresión si no se transforma."),
        (("No rellenar huecos con ceros ni promedios. ", "b"), "Trabajar con las filas completas de cada análisis e informar su n."),
        (("Estratificar. ", "b"), "Mezclar basaltos con peridotitas puede inventar o esconder relaciones; conviene repetir cada "
         "regresión por tipo de roca o por grupo tectónico."),
    ])
    I.p("Pares sugeridos por tipo de modelo: lineal simple, MgO frente a Ni; potencial o log-log, Zr frente a Hf y Nb frente a "
        "Ta; múltiple, los óxidos mayores.")
    I.salto()


def parte_tecnica(I):
    I.h("Parte III. Cómo se construyó: documentación técnica", 1)
    I.p("Esta parte explica cómo se hizo cada producto del proyecto, con qué herramientas y por qué, para que cualquier "
        "integrante (o el docente) pueda entenderlo, reproducirlo y modificarlo. Todo el código está en el repositorio.")

    I.h("11. Arquitectura general", 2)
    I.p("El proyecto es una cadena de pasos donde cada uno lee el resultado del anterior. No hay servidor ni base de datos "
        "en línea: todo se calcula en la computadora y el resultado se publica como archivos estáticos.")
    I.codigo("""
EarthChem PetDB (2 CSV crudos, ~600 MB)
        │  scripts/limpiar.py  (Python)
        ▼
data/earthchem_limpio.csv.gz  ── base limpia: 159.760 × 75
        ├─ scripts/tabla_variables.py  → tabla de variables (Excel + JSON para la web)
        ├─ scripts/datos_dashboard.py  → resumen.json (cifras, mapa y TAS de la web)
        ├─ scripts/mapa.py             → capas GIS (GeoPackage) + mapa PNG
        │     └─ scripts/proyecto_qgis.py (PyQGIS) → proyecto QGIS + mapa A4
        ├─ R/01..03_*.R                → TDF, gráficas, indicadores (PNG + CSV + Excel)
        └─ scripts/informe.py          → este informe (Word + PDF)
                    │
                    ▼
docs/  (la web estática) ── git push ──► GitHub Pages
""")

    I.h("12. Herramientas y librerías", 2)
    I.tabla([
        ["Herramienta", "Versión", "Para qué se usó"],
        ["Python", "3.13", "limpieza, tabla de variables, mapa, agregados de la web, informe"],
        ["pandas / numpy", "2.3 / 2.4", "lectura y resumen de la base limpia"],
        ["csv, gzip, re (biblioteca estándar)", "—", "lectura por streaming del CSV crudo de 2.318 columnas"],
        ["openpyxl", "3.1", "exportar la tabla de variables a Excel"],
        ["geopandas, shapely, pyproj, pyogrio", "1.1 / 2.1 / 3.7 / 0.12", "capas GIS, reproyección a Robinson, GeoPackage"],
        ["matplotlib + matplotlib-scalebar", "3.10 / 0.9", "mapa de ubicación con escala, norte y grilla"],
        ["QGIS (PyQGIS)", "3.44", "proyecto .qgz y composición de impresión A4"],
        ["Natural Earth 1:110m", "—", "límites de países (dominio público)"],
        ["R", "4.6.1", "estadística descriptiva (exigencia de la asignatura)"],
        ["e1071 / openxlsx (R)", "1.7 / 4.2", "asimetría y curtosis; exportar TDF a Excel"],
        ["HTML5, CSS3, JavaScript", "—", "la página web, sin framework"],
        ["Plotly.js (CDN)", "2.35.2", "gráficos interactivos: mapa, barras, diagrama TAS"],
        ["Google Fonts", "—", "Fraunces, Instrument Sans y JetBrains Mono"],
        ["Pillow", "12.2", "recorte de fotos, optimización de imágenes, favicon"],
        ["gflow-cli (Google Flow)", "0.83", "imágenes ilustrativas generadas con IA"],
        ["python-docx + Microsoft Word", "1.2", "este informe y su versión PDF"],
        ["Git, GitHub, GitHub Pages, gh CLI", "—", "control de versiones y publicación de la web"],
    ], [5.2, 2.6, 8.5])

    I.h("13. Limpieza de la base (scripts/limpiar.py)", 2)
    I.p("El CSV crudo tiene 2.318 columnas y celdas con saltos de línea, así que no se carga entero con pandas: se recorre "
        "fila por fila con el módulo csv de la biblioteca estándar, en dos pasadas. La primera cuenta cuántos valores "
        "numéricos tiene cada columna química y comprueba que el nombre de cada muestra coincida con el del archivo de "
        "metadatos (si el orden se rompe, el script se detiene). La segunda escribe la base limpia comprimida con gzip.")
    I.p("Cada regla de limpieza es una función pequeña y probada: el bloque principal tiene asserts que verifican, por "
        "ejemplo, que «39.244°N» se convierta en 39,244, que una latitud mayor que 90 se intercambie con la longitud y que "
        "un ambiente combinado como «FRACTURE ZONE,SPREADING CENTER» caiga en el grupo correcto:")
    I.codigo("""
def grupo_tectonico(ambiente):
    \"\"\"Valores combinados toman el grupo común; si difieren, 'mixto'.\"\"\"
    if not ambiente:
        return ""
    grupos = {g for a in ambiente.upper().split(",")
                for g, miembros in GRUPO_TECTONICO.items() if a.strip() in miembros}
    return grupos.pop() if len(grupos) == 1 else "mixto"
""")
    I.p("La tabla de variables (", ("scripts/tabla_variables.py", "c"), ") clasifica cada columna y toma como «instrumento» el "
        "método analítico más frecuente de cada variable, leído de las columnas Method del archivo crudo (por ejemplo, el "
        "SiO2 se midió con fluorescencia de rayos X en el 64 % de los casos).")

    I.h("14. Mapa de ubicación (Python y QGIS)", 2)
    I.p(("scripts/mapa.py", "c"), " convierte las coordenadas en puntos con geopandas, descarga los límites de países de "
        "Natural Earth y genera una grilla de 30° y el contorno del globo. Todo se guarda en un GeoPackage (",
        ("data/gis/mapa_ubicacion.gpkg", "c"), ") en WGS84, con nombres de campo cortos para la tabla de atributos de QGIS. "
        "Luego dibuja el mapa en proyección Robinson (ESRI:54030) con matplotlib: título, leyenda con el número de "
        "análisis por ambiente tectónico, escala gráfica, flecha de norte, coordenadas y fuente.")
    I.p(("scripts/proyecto_qgis.py", "c"), " se ejecuta con el Python que trae QGIS (PyQGIS). Carga las mismas capas, aplica "
        "una simbología categorizada con los mismos colores, arma una composición de impresión A4 y la exporta a PNG. El "
        "proyecto se guarda con rutas relativas, así que el .qgz funciona en cualquier computadora si está junto al .gpkg. "
        "Un detalle técnico: QGIS no logra ubicar etiquetas junto a los polos en la proyección Robinson, por eso las "
        "longitudes se rotulan sobre el ecuador.")
    I.imagen(RAIZ / "docs" / "img" / "mapa_qgis.png", 16, "Composición A4 generada con PyQGIS.")

    I.h("15. Estadística en R (carpeta R/)", 2)
    I.p(("R/funciones.R", "c"), " reúne las funciones comunes; los scripts ", ("01_cualitativas.R", "c"), ", ",
        ("02_discretas.R", "c"), " y ", ("03_continuas.R", "c"), " las aplican. La función central arma la TDF; sin "
        "argumentos reproduce exactamente la plantilla del curso, y con un rango típico agrega las clases abiertas:")
    I.codigo("""
tdf_def  <- tdf_continua(x, redondear = FALSE)       # plantilla: a = R / k
lim      <- limites_iqr(x)                           # Q1 - 1,5·IR ; Q3 + 1,5·IR
tdf_fin  <- tdf_continua(x, lim = lim)               # Sturges en el rango típico + clases abiertas
error_agrupacion(x, tdf_fin)                         # media y desviación agrupadas vs. totales
""")
    I.p("Cada script produce tres salidas: las gráficas en PNG (", ("docs/img/r/", "c"), "), las tablas en CSV para la web (",
        ("docs/data/tdf/", "c"), ") y un Excel con todas las tablas (", ("R/salidas/", "c"), "). Un problema encontrado: en "
        "R 4.6 para Windows, read.csv sobre el archivo .gz se cerraba con un error de segmentación; se resolvió abriéndolo "
        "con gzfile y codificación UTF-8 explícita (función leer_base).")

    I.h("16. La página web", 2)
    I.h("16.1 Estructura y tecnología", 3)
    I.p("La web es un sitio estático: archivos HTML, CSS y JavaScript que el navegador lee directamente, sin framework ni "
        "paso de compilación. Se eligió así para que GitHub Pages la sirva gratis tal cual y para que cualquier integrante "
        "pueda editarla con un editor de texto. Tiene una página por unidad del curso, que comparten un archivo de estilos "
        "y uno de código:")
    I.tabla([
        ["Archivo", "Contenido"],
        ["docs/index.html", "portada, introducción, problema, mapa (con pestañas QGIS / Python / interactivo), objetivos"],
        ["docs/descriptiva/index.html", "Unidad 1: población y muestra, tabla de variables, TDF y gráficas de R"],
        ["docs/datos/index.html", "origen de los datos, preparación y descargas"],
        ["docs/equipo/index.html", "integrantes y docente"],
        ["docs/assets/estilo.css", "todo el diseño: colores, tipografía, grillas, animaciones, versión móvil"],
        ["docs/assets/app.js", "código común: navegación, animaciones, pestañas, lectura de CSV, estilo de gráficos"],
        ["docs/data/", "resumen.json, variables.json y las TDF en CSV"],
        ["docs/img/", "fotos, mapas, gráficas de R e imágenes ilustrativas"],
    ], [5.5, 10.8])
    I.h("16.2 Cómo carga los datos", 3)
    I.p("La página no contiene los datos escritos en el HTML: los pide al cargarse con fetch y los dibuja. La tabla de "
        "variables viene de variables.json; los gráficos interactivos, de resumen.json (254 KB con agregados y muestras de "
        "6.000 puntos, en lugar de la base completa de 26 MB); y las TDF, de los CSV que exporta R. Para leer esos CSV sin "
        "librerías se escribió un lector de pocas líneas:")
    I.codigo("""
async function leerCSV(url) {
  const texto = (await (await fetch(url)).text()).trim();
  const campos = linea => [...linea.matchAll(/("([^"]*)"|[^,]*)(,|$)/g)]
                            .map(m => m[2] ?? m[1]).slice(0, -1);
  const [cab, ...filas] = texto.split(/\\r?\\n/).map(campos);
  return filas.map(f => Object.fromEntries(cab.map((c, i) => [c, f[i]])));
}
""")
    I.p("Por eso la web debe verse a través de un servidor (GitHub Pages, o localmente con ", ("python -m http.server", "c"),
        "): abriendo el archivo con doble clic, el navegador bloquea esas lecturas.")
    I.h("16.3 Componentes y comportamiento", 3)
    I.lista([
        (("Navegación entre páginas: ", "b"), "un menú fijo marca la página actual (atributo data-pagina del body); las "
         "unidades aún no hechas aparecen como «pronto»."),
        (("Aparición al hacer scroll: ", "b"), "los bloques con clase rev aparecen con un desplazamiento suave usando "
         "IntersectionObserver, una API del navegador que avisa cuando un elemento entra en pantalla."),
        (("Cifras animadas: ", "b"), "las cifras de la portada cuentan hacia arriba con requestAnimationFrame."),
        (("Pestañas: ", "b"), "una función genérica muestra un panel y oculta los demás (mapa, variables continuas)."),
        (("Tabla de variables filtrable: ", "b"), "filtros por tipo y un buscador que filtra mientras se escribe."),
        (("Gráficos interactivos: ", "b"), "Plotly.js dibuja el mapa mundial, las barras de ambientes tectónicos y el "
         "diagrama TAS, con información al pasar el cursor."),
    ])
    I.h("16.4 Diseño", 3)
    I.p("La dirección visual es la de una revista científica sobre un tema oscuro «basalto»: fondo casi negro cálido, texto "
        "color hueso y un acento color magma. Los colores se definen una sola vez como variables CSS y el resto del código "
        "las usa por nombre, de modo que cambiar un color se hace en un solo lugar:")
    I.codigo("""
:root {
  --page: #0f0e0c;  --surface-1: #171513;  --text-primary: #f0ece4;
  --magma: #ff6b2c; --oro: #e9b44c;
  --series-1: #3987e5; --series-2: #d95926; --series-3: #199e70;
  --serif: "Fraunces", Georgia, serif; --sans: "Instrument Sans", sans-serif;
}
""")
    I.p("Los colores de los datos (azul, naranja y verde agua, más una rampa azul para las variables ordinales) son una "
        "paleta validada para personas con daltonismo, y son los mismos en la web, en el mapa de Python y en el de QGIS. "
        "Nunca se usan más de tres colores de categoría en un mismo gráfico. La tipografía combina Fraunces para títulos, "
        "Instrument Sans para el texto y JetBrains Mono para cifras y fórmulas.")
    I.h("16.5 Accesibilidad, celulares y rendimiento", 3)
    I.lista([
        "Todas las imágenes tienen texto alternativo; las pestañas y botones de filtro usan atributos ARIA.",
        "Si el sistema pide reducir el movimiento (prefers-reduced-motion), se desactivan las animaciones.",
        "Debajo de 900 px de ancho las grillas pasan a una o dos columnas y el menú se desplaza horizontalmente; se "
        "verificó que ninguna página se desborde en una pantalla de celular de 390 px.",
        "Las imágenes fuera de pantalla se cargan recién al acercarse (loading=\"lazy\") y los scripts se cargan con defer.",
        "Las imágenes generadas se comprimieron a JPEG (unos 170–335 KB cada una).",
    ])
    I.h("16.6 Imágenes", 3)
    I.p("Las fotos del equipo son reales: se recortaron a 400 × 400 px con Pillow. Las imágenes ilustrativas (lámina "
        "delgada de basalto, Andes ecuatorianos, buque de perforación, testigos de roca) se generaron con Google Flow desde "
        "la terminal con gflow-cli, por ejemplo:")
    I.codigo('gflow image t2i "Macro photograph of a basalt thin section under cross-polarized light..." \\\n'
             '      --model nano2 --aspect 16:9 -n 1 --out <carpeta>')
    I.p("La web indica en los pies y en el pie de página que esas imágenes son ilustrativas y fueron generadas con IA. El "
        "favicon (el logo «Pt») se dibujó en SVG y se exportó a ICO y PNG con Pillow.")

    I.h("17. Publicación y control de versiones", 2)
    I.p("El proyecto vive en un repositorio Git publicado en GitHub. Cada cambio se registra como un commit con un mensaje "
        "que explica qué se hizo, así que el historial documenta cómo evolucionó el trabajo. GitHub Pages está configurado "
        "para publicar la carpeta docs/ de la rama main: cada git push que modifica esa carpeta actualiza la web en uno a "
        "tres minutos. El repositorio es público porque GitHub Pages gratuito no funciona con repositorios privados; los "
        "datos tienen licencia CC BY 4.0, que permite publicarlos citando la fuente.")
    I.p("El repositorio pertenece a la cuenta de Cristoffer Rumiguano (crisrumiguano-art), a quien se transfirió el 15 de "
        "octubre de 2026; Dalember (cuenta DalemberV) colabora con permiso de escritura. Cualquier integrante con acceso puede proponer "
        "cambios con un pull request.")
    I.p("Límites a tener en cuenta: GitHub rechaza archivos de más de 100 MB, por eso los datos crudos (~600 MB) no están en "
        "el repositorio. Los archivos más pesados que sí están son las capas GIS (~35 MB) y la base limpia (~26 MB).")

    I.h("18. Documentación del proyecto", 2)
    I.lista([
        (("README.md: ", "b"), "la portada del repositorio en GitHub, con enlaces a todo."),
        (("AGENTS.md: ", "b"), "reglas y decisiones del proyecto (requisitos del curso, decisiones de limpieza, trampas de "
         "los datos, diseño). Lo leen los asistentes de programación con IA (Claude Code, Gemini, Codex); CLAUDE.md y "
         "GEMINI.md lo importan."),
        (("openwiki/: ", "b"), "documentación explicativa en formato OpenWiki: base de datos, cómo regenerar y publicar, y "
         "estructura de la web."),
        (("R/analisis_agrupacion_outliers.md: ", "b"), "el razonamiento de la estadística descriptiva."),
        (("Este informe: ", "b"), "se genera con scripts/informe.py, así que se actualiza solo si cambian los datos."),
    ])

    I.h("19. Cómo reproducir todo", 2)
    I.p("Los dos CSV crudos están respaldados en Google Drive (carpeta 10-26 GDRIVE GEOQUIMICA ROCAS › datos_crudos_earthchem). "
        "En la computadora de Dalember, data/raw/crudos es un enlace (junction) a esa carpeta, y los scripts la usan por "
        "defecto. En otra computadora, basta con copiar ahí los CSV o pasar su carpeta como argumento. Desde la raíz del "
        "repositorio y en este orden:")
    I.codigo("""
pip install -r requirements.txt
python -I scripts/limpiar.py            # usa data/raw/crudos
python -I scripts/datos_dashboard.py
python scripts/tabla_variables.py      # usa data/raw/crudos
python -I scripts/mapa.py
"C:\\Program Files\\QGIS 3.44.8\\bin\\python-qgis-ltr.bat" scripts/proyecto_qgis.py
Rscript R/01_cualitativas.R
Rscript R/02_discretas.R
Rscript R/03_continuas.R
python scripts/informe.py
git add -A && git commit -m "..." && git push
""")
    I.p("Para revisar la web antes de publicarla: ", ("python -m http.server 8765 -d docs", "c"), " y abrir "
        "http://localhost:8765/ en el navegador.")

    I.h("20. Limitaciones técnicas y mantenimiento", 2)
    I.lista([
        "El menú de navegación está repetido en cada página: al agregar una página (por ejemplo, Inferencial) hay que "
        "actualizarlo en todas.",
        "Los textos de lectura de las gráficas en la página descriptiva están escritos a mano con las cifras actuales; si "
        "cambian los datos, hay que revisarlos.",
        "El informe usa Microsoft Word para generar el PDF y actualizar el índice; sin Word queda solo el .docx.",
        "mapa.py necesita internet para descargar los límites de Natural Earth.",
        "gflow-cli automatiza la interfaz web de Google Flow y puede dejar de funcionar si Google la cambia.",
    ])
    I.salto()


def anexos(I, df):
    I.h("Anexo A. Diccionario de variables", 1)
    I.tabla([["Columna", "Tipo", "% con dato"]]
            + [[c, TIPOS.get(c, "continua"), f"{100 * df[c].notna().mean():.1f}".replace(".", ",")] for c in df.columns],
            [7.5, 5.5, 3])
    I.h("Anexo B. Estructura del repositorio", 1)
    I.codigo("""
data/            base limpia, tabla de variables, capas GIS y proyecto QGIS
  raw/           consulta de descarga y términos de uso
scripts/         Python: limpieza, tabla de variables, agregados web, mapa, QGIS, informe
R/               estadística descriptiva (funciones + 3 scripts) y su análisis
  salidas/       Excel con todas las TDF
docs/            la web publicada en GitHub Pages
  assets/        estilo.css y app.js
  descriptiva/ datos/ equipo/   páginas
  data/          resumen.json, variables.json, tdf/*.csv
  img/           fotos, mapas, gráficas de R (img/r/) e imágenes generadas
openwiki/        documentación explicativa
AGENTS.md        reglas del proyecto (CLAUDE.md y GEMINI.md lo importan)
README.md        portada del repositorio
""")


def a_pdf():
    """Abre el .docx en Word, actualiza índice y campos, guarda y exporta a PDF."""
    ps = f"""
$w = New-Object -ComObject Word.Application; $w.Visible = $false
try {{
  $d = $w.Documents.Open('{DOCX}')
  foreach ($t in $d.TablesOfContents) {{ $t.Update() }}
  $d.Fields.Update() | Out-Null
  $d.Save()
  $d.ExportAsFixedFormat('{PDF}', 17)
  $d.Close()
}} finally {{ $w.Quit() }}
"""
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print("PDF:", PDF if r.returncode == 0 and PDF.exists() else f"no generado ({r.stderr.strip()[:200]})")


def main():
    df = pd.read_csv(RAIZ / "data" / "earthchem_limpio.csv.gz", low_memory=False)
    I = Informe()
    portada(I)
    guia_equipo(I, df)
    parte_base(I, df)
    parte_descriptiva(I, df)
    parte_tecnica(I)
    anexos(I, df)
    I.d.core_properties.title = "Geoquímica de rocas a escala mundial — Informe del proyecto"
    I.d.core_properties.author = "Rumiguano, Fernández, Bustamante, Haro"
    I.d.save(DOCX)
    print("Word:", DOCX)
    a_pdf()


if __name__ == "__main__":
    main()
